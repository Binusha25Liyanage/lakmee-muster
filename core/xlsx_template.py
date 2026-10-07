"""Excel layout templates: learn the structure of an example Excel output and reuse it.

extract(path)  reads an example .xlsx (e.g. Transaction_2026_10_03.xlsx or Attendance_Week1_...xlsx) and returns a template:
               title / sub-title rows, header names and order, colours, fonts, banding, widths, day-group row.
apply(spec, tpl) re-shapes a workbook spec made by a module (canonical fields) so it follows that template.
Two kinds: "list" (one table, one row per record) and "grid" (day groups with a Check in / Check out pair under each day).
"""
import datetime
import re
import uuid

import openpyxl
from openpyxl.utils import get_column_letter

SYN = {
    "id": ["employee id", "emp id", "id", "emp no", "employee no", "staff id", "employee number"],
    "name": ["first name", "name", "employee", "employee name", "full name", "staff name"],
    "dept": ["department", "dept", "section"],
    "date": ["date", "day"],
    "time": ["time", "punch time", "clock"],
    "state": ["punch state", "state", "punch", "status"],
    "in": ["check in", "check-in", "clock in", "in", "arrival", "first check-in"],
    "out": ["check out", "check-out", "clock out", "out", "last check-out"],
}
FIELD_LABEL = {"id": "Employee ID", "name": "First Name", "dept": "Department", "date": "Date", "time": "Time",
               "state": "Punch State", "in": "Check in", "out": "Check out"}
DEFAULT_LOOK = {
    "title": {"fill": "4A4743", "color": "FFFFFF", "size": 14, "bold": True, "align": "center", "height": 24},
    "subtitle": {"bold": True, "align": "left"},
    "group": {"fill": "4A4743", "color": "FFFFFF", "bold": True, "align": "center"},
    "header": {"fill": "4A4743", "color": "FFFFFF", "bold": True, "align": "left"},
    "body": {"font": "Calibri", "size": 11, "band": 0, "band_color": "EFE9F5", "border": "CBBFAE"},
}


def field_of(label):
    t = re.sub(r"\s+", " ", str(label or "").strip().lower())
    for f, names in SYN.items():
        if t in names:
            return f
    return None


def _rgb(color, default=None):
    try:
        if color is not None and color.type == "rgb" and isinstance(color.rgb, str):
            return color.rgb[-6:].upper()
    except Exception:
        pass
    return default


def _fill(cell):
    try:
        if cell.fill and cell.fill.fill_type == "solid":
            return _rgb(cell.fill.fgColor)
    except Exception:
        pass
    return None


def _cell_look(cell, default):
    d = dict(default)
    f = _fill(cell)
    d["fill"] = f
    d["color"] = _rgb(cell.font.color, "000000") if cell.font else "000000"
    d["bold"] = bool(cell.font and cell.font.b)
    if cell.font and cell.font.sz:
        d["size"] = float(cell.font.sz)
    d["align"] = (cell.alignment.horizontal if cell.alignment and cell.alignment.horizontal else "left")
    return d


def _widths(ws):
    out = {}
    for dim in ws.column_dimensions.values():
        if dim.width and dim.min:
            for c in range(dim.min, (dim.max or dim.min) + 1):
                out[c] = float(dim.width)
    return out


def _date_pattern(text):
    """'Sep 1' -> 'MMM D', '1 Sep' -> 'D MMM', '2026-09-01' -> 'YYYY-MM-DD'."""
    t = str(text or "").strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", t):
        return "YYYY-MM-DD"
    if re.fullmatch(r"[A-Za-z]{3,9}\.? \d{1,2}", t):
        return "MMM D"
    if re.fullmatch(r"\d{1,2}[ -][A-Za-z]{3,9}", t):
        return "D MMM"
    if re.fullmatch(r"\d{1,2}/\d{1,2}", t):
        return "D/M"
    return "MMM D"


def fmt_date(iso, pattern):
    d = datetime.date.fromisoformat(iso)
    return {"MMM D": f"{d.strftime('%b')} {d.day}", "D MMM": f"{d.day} {d.strftime('%b')}", "YYYY-MM-DD": d.isoformat(),
            "D/M": f"{d.day}/{d.month}"}.get(pattern, f"{d.strftime('%b')} {d.day}")


def _text_cells(ws, r, maxc):
    return [(c, ws.cell(r, c).value) for c in range(1, maxc + 1) if ws.cell(r, c).value not in (None, "")]


def extract(path, name=None):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    maxc = ws.max_column
    merged = list(ws.merged_cells.ranges)
    # header row = first row with >= 3 text cells whose next row also has content
    def is_group_row(r):          # a row of merged day labels (Sep 1 | Sep 2 ...) sits ABOVE the header row
        return sum(1 for m in merged if m.min_row == m.max_row == r and m.max_col > m.min_col and ws.cell(r, m.min_col).value) >= 2
    hr = next((r for r in range(1, min(ws.max_row, 12) + 1)
               if len(_text_cells(ws, r, maxc)) >= 3 and _text_cells(ws, r + 1, maxc) and not is_group_row(r)), None)
    if hr is None:
        raise ValueError("Could not find the header row (a row with at least 3 column names above the data).")
    groups = [m for m in merged if m.min_row == m.max_row == hr - 1 and m.max_col > m.min_col and ws.cell(m.min_row, m.min_col).value]
    kind = "grid" if len(groups) >= 2 else "list"
    heads = [(c, ws.cell(hr, c).value) for c in range(1, maxc + 1) if ws.cell(hr, c).value not in (None, "")]
    widths = _widths(ws)
    tpl = {"id": "xl-" + uuid.uuid4().hex[:8], "name": name or "Excel layout", "kind": kind, "builtin": False,
           "sheet_name": ws.title, "look": {k: dict(v) for k, v in DEFAULT_LOOK.items()}, "notes": []}
    # ---- title / subtitle rows
    title_r = 1 if hr > 1 and _text_cells(ws, 1, maxc) else None
    if title_r:
        t = ws.cell(1, 1)
        span = next((m.max_col for m in merged if m.min_row == 1 == m.max_row and m.min_col == 1), 1)
        tpl["title"] = re.sub(r"\d{4}-\d{2}-\d{2}", "{date}", str(t.value))
        tpl["title_span"] = "full" if span >= maxc else span
        lk = _cell_look(t, DEFAULT_LOOK["title"])
        lk["height"] = ws.row_dimensions[1].height or 22
        tpl["look"]["title"] = lk
    else:
        tpl["title"] = None
    sub_r = next((r for r in range(2, hr) if kind == "list" and ws.cell(r, 1).value), None)
    if sub_r:
        c = ws.cell(sub_r, 1)
        tpl["subtitle"] = re.sub(r"\d{4}-\d{2}-\d{2}", "{date}", str(c.value))
        tpl["look"]["subtitle"] = _cell_look(c, DEFAULT_LOOK["subtitle"])
    else:
        tpl["subtitle"] = None
    # ---- header + body look
    tpl["look"]["header"] = _cell_look(ws.cell(hr, heads[0][0]), DEFAULT_LOOK["header"])
    b1, b2 = ws.cell(hr + 1, 1), ws.cell(hr + 2, 1)
    f1, f2 = _fill(b1), _fill(b2)
    body = dict(DEFAULT_LOOK["body"])
    body["font"] = b1.font.name or "Calibri"
    body["size"] = float(b1.font.sz or 11)
    body["border"] = _rgb(b1.border.left.color, "CBBFAE") if b1.border and b1.border.left and b1.border.left.style else None
    if f1 or f2:
        body["band"], body["band_color"] = (0, f1) if f1 else (1, f2)
    else:
        body["band_color"] = None
    tpl["look"]["body"] = body
    # ---- columns
    if kind == "list":
        cols = []
        for c, h in heads:
            cols.append({"header": str(h), "field": field_of(h), "width": widths.get(c), "align": (ws.cell(hr + 1, c).alignment.horizontal or "left")[0]})
        tpl["columns"] = cols
        unknown = [x["header"] for x in cols if not x["field"]]
        if unknown:
            tpl["notes"].append("Columns the app has no data for will stay empty: " + ", ".join(unknown))
    else:
        first = min(m.min_col for m in groups)
        lead = []
        for c, h in heads:
            if c < first:
                lead.append({"header": str(h), "field": field_of(h), "width": widths.get(c), "align": (ws.cell(hr + 1, c).alignment.horizontal or "left")[0]})
        g0 = sorted(groups, key=lambda m: m.min_col)[0]
        pair = [str(ws.cell(hr, c).value or "") for c in range(g0.min_col, g0.max_col + 1)]
        tpl["columns"], tpl["pair"] = lead, pair
        tpl["pair_fields"] = [field_of(p) or ("in" if i == 0 else "out") for i, p in enumerate(pair)]
        tpl["pair_width"] = widths.get(g0.min_col)
        tpl["pair_align"] = (ws.cell(hr + 1, g0.min_col).alignment.horizontal or "center")[0]
        tpl["group_format"] = _date_pattern(ws.cell(g0.min_row, g0.min_col).value)
        tpl["look"]["group"] = _cell_look(ws.cell(g0.min_row, g0.min_col), DEFAULT_LOOK["group"])
    return tpl


def summary(t):
    cols = [c["header"] for c in t["columns"]]
    if t["kind"] == "grid":
        return "Weekly grid: " + ", ".join(cols) + " + a day group of [" + " | ".join(t["pair"]) + "] per day"
    return "List: " + ", ".join(cols)


# ------------------------------------------------------------------ apply
def apply(spec, tpl):
    """spec = canonical workbook sheet from a module (spec['fields']), tpl = template. Returns a new spec."""
    if not tpl or tpl.get("kind") != spec.get("kind"):
        return spec
    s = dict(spec)
    look = {k: {**DEFAULT_LOOK[k], **(tpl.get("look", {}).get(k) or {})} for k in DEFAULT_LOOK}
    s["look"] = look
    fields = spec["fields"]
    if spec["kind"] == "list":
        iso = spec.get("date") or ""
        cols = tpl["columns"]
        idx = {f: i for i, f in enumerate(fields)}
        rows = []
        for r in spec["rows"]:
            line = []
            for c in cols:
                f = c["field"]
                line.append(iso if f == "date" else (r[idx[f]] if f in idx else ""))
            rows.append(line)
        s.update(headers=[c["header"] for c in cols], rows=rows, aligns=[c["align"] for c in cols],
                 widths={i + 1: (c["width"] or 14) for i, c in enumerate(cols)},
                 title=(tpl.get("title") or "").replace("{date}", iso) if tpl.get("title") else None,
                 subtitle=(tpl.get("subtitle") or "").replace("{date}", iso) if tpl.get("subtitle") else None,
                 title_cols=len(cols) if tpl.get("title_span") == "full" else min(tpl.get("title_span") or 1, len(cols)))
        s["name"] = tpl.get("sheet_name") or spec["name"]
    else:
        lead = tpl["columns"]
        idx = {f: i for i, f in enumerate(fields)}
        days = spec["days"]
        pair, pf = tpl["pair"], tpl["pair_fields"]
        heads, group, rows = [c["header"] for c in lead], [], []
        col = len(lead) + 1
        for d in days:
            group.append((fmt_date(d, tpl.get("group_format", "MMM D")), col, col + len(pair) - 1))
            heads += pair
            col += len(pair)
        n_lead = len(fields)
        for r in spec["rows"]:
            line = [r[idx[c["field"]]] if c["field"] in idx else "" for c in lead]
            for k in range(len(days)):
                for j, f in enumerate(pf):
                    base = n_lead + k * 2
                    line.append(r[base + (0 if f == "in" else 1)])
            rows.append(line)
        widths = {i + 1: (c["width"] or 16) for i, c in enumerate(lead)}
        for c in range(len(lead) + 1, len(heads) + 1):
            widths[c] = tpl.get("pair_width") or 11
        s.update(headers=heads, group=group, rows=rows, widths=widths,
                 aligns=[c["align"] for c in lead] + [tpl.get("pair_align", "c")] * (len(heads) - len(lead)),
                 title=tpl.get("title"), title_cols=(len(lead) if tpl.get("title_span") == "full" else min(tpl.get("title_span") or 1, len(heads))))
        s["flat_headers"] = [c["header"] for c in lead] + [f"{g[0]} {p}" for g in group for p in pair]
        s["headers_preview"] = s["flat_headers"]
    return s
