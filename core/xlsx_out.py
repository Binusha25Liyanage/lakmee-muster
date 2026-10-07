"""Writes the Excel (.xlsx) outputs. Looks are copied from the example files:
  style "daily"  - Transaction_YYYY_MM_DD.xlsx   (title row, date row, header row, banded list)
  style "weekly" - Attendance_WeekN_...xlsx      (title, day group row, Check in / Check out columns)
A workbook can hold several sheets (the monthly output = one sheet per week).
A sheet spec may carry "look" (colours, fonts, banding) learned from an uploaded Excel layout (see xlsx_template.py).
"""
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LOOK = {
    "title": {"fill": "4A4743", "color": "FFFFFF", "size": None, "bold": True, "align": "center", "height": None},
    "subtitle": {"bold": True, "align": "left"},
    "group": {"fill": "4A4743", "color": "FFFFFF", "bold": True, "align": "center"},
    "header": {"fill": None, "color": "000000", "bold": True, "align": "center"},
    "body": {"font": "Calibri", "size": 11, "band": 0, "band_color": None, "border": "CBBFAE"},
}


def _fill(hexcode):
    return PatternFill("solid", start_color="FF" + hexcode, end_color="FF" + hexcode) if hexcode else PatternFill()


def _sheet_name(name, used):
    n = re.sub(r"[\\/*?:\[\]]", "_", name)[:31] or "Sheet"
    base, i = n, 2
    while n.lower() in used:
        n = f"{base[:28]}_{i}"; i += 1
    used.add(n.lower())
    return n


def write_workbook(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for spec in sheets:
        ws = wb.create_sheet(_sheet_name(spec["name"], used))
        (_daily if spec["style"] == "daily" else _weekly)(ws, spec)
    wb.save(path)
    return path


def _look(spec):
    """spec['look'] wins; otherwise the built-in look of the example sheets (daily / weekly)."""
    if spec.get("look"):
        return spec["look"]
    base = {k: dict(v) for k, v in LOOK.items()}
    base["title"]["size"] = 14 if spec["style"] == "daily" else 13
    base["title"]["height"] = 24 if spec["style"] == "daily" else 22
    if spec["style"] == "daily":
        base["header"] = {"fill": spec.get("head_fill", "4A4743"), "color": "FFFFFF", "bold": True, "align": "left"}
        base["body"]["band"], base["body"]["band_color"] = spec.get("band", 0), spec.get("band_color")
    else:
        base["header"]["fill"] = "F0EDE9"
        base["body"]["band"], base["body"]["band_color"] = spec.get("band", 1), spec.get("band_color")
    return base


def _align(a):
    a = (a or "left")
    return {"l": "left", "c": "center", "r": "right"}.get(a, a)


def _font(l, **kw):
    return Font(name=kw.get("name"), size=l.get("size"), bold=bool(l.get("bold")), color="FF" + (l.get("color") or "000000"))


def _box(color):
    if not color:
        return Border()
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def _widths(ws, spec):
    for c, w in spec["widths"].items():
        ws.column_dimensions[get_column_letter(c)].width = w


def _body(ws, spec, first_row, look):
    b = look["body"]
    ncol = len(spec["headers"])
    for i, row in enumerate(spec["rows"]):
        r = first_row + i
        banded = bool(b.get("band_color")) and i % 2 == b.get("band", 0)
        for c in range(ncol):
            v = row[c] if c < len(row) else None
            cell = ws.cell(r, c + 1, v if v not in ("",) else None)
            cell.font = Font(name=b.get("font") or "Calibri", size=b.get("size") or 11)
            cell.alignment = Alignment(horizontal=_align(spec["aligns"][c]), vertical="center")
            cell.border = _box(b.get("border"))
            if banded:
                cell.fill = _fill(b["band_color"])


def _title(ws, spec, look, ncols):
    if not spec.get("title"):
        return 1
    t = look["title"]
    span = spec.get("title_cols") or ncols
    if span > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(1, 1, spec["title"])
    c.font = Font(size=t.get("size") or 14, bold=bool(t.get("bold")), color="FF" + (t.get("color") or "FFFFFF"))
    if t.get("fill"):
        c.fill = _fill(t["fill"])
    c.alignment = Alignment(horizontal=_align(t.get("align")), vertical="center")
    if t.get("height"):
        ws.row_dimensions[1].height = t["height"]
    return 2


def _head(ws, spec, look, row):
    h = look["header"]
    for c, text in enumerate(spec["headers"], start=1):
        cell = ws.cell(row, c, text)
        cell.font = Font(bold=bool(h.get("bold")), color="FF" + (h.get("color") or "000000"))
        if h.get("fill"):
            cell.fill = _fill(h["fill"])
        cell.alignment = Alignment(horizontal=_align(h.get("align")), vertical="center")
        cell.border = _box(look["body"].get("border") or "CBBFAE")


def _daily(ws, spec):
    look = _look(spec)
    n = len(spec["headers"])
    r = _title(ws, spec, look, n)
    if spec.get("subtitle"):
        s = look["subtitle"]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
        d = ws.cell(r, 1, spec["subtitle"])
        d.font = Font(bold=bool(s.get("bold"))); d.alignment = Alignment(horizontal=_align(s.get("align")))
        r += 1
    _head(ws, spec, look, r)
    _body(ws, spec, r + 1, look)
    _widths(ws, spec)


def _weekly(ws, spec):
    look = _look(spec)
    n = len(spec["headers"])
    _title(ws, spec, look, n)
    g = look["group"]
    for text, c1, c2 in spec["group"]:
        if c2 > c1:
            ws.merge_cells(start_row=2, start_column=c1, end_row=2, end_column=c2)
        cell = ws.cell(2, c1, text)
        cell.font = Font(bold=bool(g.get("bold")), color="FF" + (g.get("color") or "FFFFFF"))
        if g.get("fill"):
            cell.fill = _fill(g["fill"])
        cell.alignment = Alignment(horizontal=_align(g.get("align")), vertical="center")
    _head(ws, spec, look, 3)
    _body(ws, spec, 4, look)
    _widths(ws, spec)
